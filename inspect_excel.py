import os
import openpyxl

base = r'C:\Users\Student\Desktop\invoice generator\prices lists'
for fn in sorted(os.listdir(base)):
    path = os.path.join(base, fn)
    if not fn.lower().endswith(('.xlsx', '.xls')):
        continue
    print(f'\nFILE: {fn}')
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    print('Sheets:', wb.sheetnames)
    for ws in wb.worksheets[:3]:
        print('Sheet:', ws.title)
        for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(12, ws.max_row), values_only=True), 1):
            print(idx, row)
        print('---')
